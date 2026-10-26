from __future__ import annotations

import unittest
from datetime import date
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from insightquest.analytics import (
    build_cohort_installation_rate,
    build_installation_periods,
    build_inventory_risk,
    build_map_clusters,
    build_purchase_install_comparison,
    build_segment_periods,
    build_shipping_lag,
)
from insightquest.exports import build_excel_workbook


class AnalyticsTests(unittest.TestCase):
    def test_map_clusters_support_numeric_device_ids(self) -> None:
        data = pd.DataFrame(
            {
                "device_name": [101, 102],
                "customer_name": [1, 2],
                "latitude": [45.501, 45.502],
                "longitude": [-73.561, -73.562],
                "installed_at": pd.to_datetime(["2025-01-01", "2025-01-02"]),
            }
        )

        result = build_map_clusters(data, 2)

        self.assertEqual(len(result), 1)
        self.assertEqual(result.iloc[0]["installs"], 2)
        self.assertEqual(result.iloc[0]["customers"], 2)
        self.assertEqual(set(result.iloc[0]["top_devices"].split(", ")), {"101", "102"})
        self.assertGreater(result.iloc[0]["radius"], 0)

    def test_map_clusters_count_installations_without_device_labels(self) -> None:
        data = pd.DataFrame(
            {
                "device_name": [None],
                "latitude": [45.5],
                "longitude": [-73.56],
                "installed_at": pd.to_datetime(["2025-01-01"]),
            }
        )

        result = build_map_clusters(data, 2)

        self.assertEqual(result.iloc[0]["installs"], 1)
        self.assertEqual(result.iloc[0]["top_devices"], "Unknown")

    def test_excel_export_sanitizes_sheet_names_and_formula_values(self) -> None:
        workbook_bytes = build_excel_workbook(
            {
                "Bad/name?": pd.DataFrame(
                    {
                        "value": ["=1+1", "ordinary", "@cmd"],
                        "infinite_value": [float("inf")] * 3,
                        "created_at": pd.to_datetime(
                            ["2025-01-01T00:00:00Z"] * 3, utc=True
                        ),
                    }
                ),
                "bad/name?": pd.DataFrame({"value": ["second sheet"]}),
            }
        )
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)

        self.assertEqual(workbook.sheetnames, ["Bad_name_", "bad_name_ (2)"])
        self.assertEqual(workbook["Bad_name_"]["A2"].value, "=1+1")
        self.assertEqual(workbook["Bad_name_"]["A2"].data_type, "s")
        self.assertEqual(workbook["Bad_name_"]["A3"].value, "ordinary")
        self.assertEqual(workbook["Bad_name_"]["A4"].value, "@cmd")
        self.assertIsNone(workbook["Bad_name_"]["B2"].value)
        self.assertEqual(
            workbook["Bad_name_"]["C2"].value,
            pd.Timestamp("2025-01-01"),
        )

    def test_purchase_install_comparison_fills_empty_periods(self) -> None:
        start = date(2025, 1, 1)
        end = date(2025, 1, 3)
        purchases = pd.DataFrame(
            {"purchase_date": ["2025-01-01"], "sales": [2.5]}
        )
        installations = pd.DataFrame({"installed_at": ["2025-01-03"]})

        result = build_purchase_install_comparison(
            purchases, installations, "Day", start, end
        )

        self.assertEqual(result["purchases"].tolist(), [2.5, 0.0, 0.0])
        self.assertEqual(result["installs"].tolist(), [0, 0, 1])

    def test_empty_installation_history_still_returns_calendar(self) -> None:
        result = build_installation_periods(
            pd.DataFrame(columns=["installed_at"]),
            "Month",
            date(2025, 1, 1),
            date(2025, 3, 31),
        )

        self.assertEqual(len(result), 3)
        self.assertEqual(result["installs"].tolist(), [0, 0, 0])

    def test_sales_periods_aggregates_weeks_from_monday_and_months(self) -> None:
        sales = pd.DataFrame(
            {
                "purchase_date": ["2025-01-05", "2025-01-06", "2025-02-01"],
                "sales": [1, 2, 4],
            }
        )
        weekly = build_purchase_install_comparison(
            sales,
            pd.DataFrame(columns=["installed_at"]),
            "Week",
            date(2025, 1, 5),
            date(2025, 2, 2),
        )
        monthly = build_purchase_install_comparison(
            sales,
            pd.DataFrame(columns=["installed_at"]),
            "Month",
            date(2025, 1, 1),
            date(2025, 2, 28),
        )

        self.assertEqual(
            weekly["period_start"].tolist(),
            [
                pd.Timestamp("2024-12-30"),
                pd.Timestamp("2025-01-06"),
                pd.Timestamp("2025-01-13"),
                pd.Timestamp("2025-01-20"),
                pd.Timestamp("2025-01-27"),
            ],
        )
        self.assertEqual(weekly["purchases"].tolist(), [1, 2, 0, 0, 4])
        self.assertEqual(monthly["purchases"].tolist(), [3, 4])

    def test_segment_periods_normalizes_unknown_and_groups_by_week(self) -> None:
        data = pd.DataFrame(
            {
                "installed_at": ["2025-01-07", "2025-01-08", "2025-01-14"],
                "segment": ["North", None, "North"],
            }
        )

        result = build_segment_periods(data, "installed_at", "segment", "Week")

        self.assertEqual(
            list(result.itertuples(index=False, name=None)),
            [
                (pd.Timestamp("2025-01-06"), "North", 1),
                (pd.Timestamp("2025-01-06"), "Unknown", 1),
                (pd.Timestamp("2025-01-13"), "North", 1),
            ],
        )

    def test_shipping_lag_reports_missing_and_negative_dates(self) -> None:
        data = pd.DataFrame(
            {
                "shipping_date": ["2025-01-01", None, "2025-01-05"],
                "installed_at": ["2025-01-03", "2025-01-04", "2025-01-04"],
            }
        )

        lags, summary = build_shipping_lag(data)

        self.assertEqual(lags["lag_days"].tolist(), [2.0])
        self.assertEqual(summary.matched_rows, 1)
        self.assertEqual(summary.missing_date_rows, 1)
        self.assertEqual(summary.negative_lag_rows, 1)
        self.assertEqual(summary.percentile_90_days, 2.0)

    def test_cohort_rate_uses_only_valid_positive_shipped_units(self) -> None:
        data = pd.DataFrame(
            {
                "shipping_date": [
                    "2025-01-01",
                    "2025-01-10",
                    "2025-02-01",
                    "2025-02-02",
                ],
                "installed_at": [
                    "2025-01-05",
                    None,
                    "2025-01-31",
                    "2025-02-03",
                ],
                "eligible_units": [2, 3, 1, 4],
            }
        )

        result = build_cohort_installation_rate(data, "eligible_units")

        self.assertEqual(result["shipped_units"].tolist(), [5, 5])
        self.assertEqual(result["installed_units"].tolist(), [2, 4])
        self.assertEqual(result["installation_rate_pct"].tolist(), [40.0, 80.0])

    def test_inventory_risk_uses_reorder_and_max_stock_boundaries(self) -> None:
        data = pd.DataFrame(
            [
                {
                    "item_name": "Low",
                    "inventory_on_hand": 10,
                    "inventory_daily_demand": 2,
                    "replenishment_days": 5,
                    "reorder_point": 5,
                    "max_stock": 15,
                },
                {
                    "item_name": "High",
                    "inventory_on_hand": 20,
                    "inventory_daily_demand": 0,
                    "replenishment_days": 2,
                    "reorder_point": 5,
                    "max_stock": 15,
                },
                {
                    "item_name": "OK",
                    "inventory_on_hand": 10,
                    "inventory_daily_demand": 1,
                    "replenishment_days": 1,
                    "reorder_point": 5,
                    "max_stock": 15,
                },
            ]
        )

        result = build_inventory_risk(data).set_index("item")

        self.assertEqual(result.loc["Low", "risk"], "Stockout risk")
        self.assertEqual(result.loc["High", "risk"], "Overstock risk")
        self.assertEqual(result.loc["OK", "risk"], "Within thresholds")


if __name__ == "__main__":
    unittest.main()
