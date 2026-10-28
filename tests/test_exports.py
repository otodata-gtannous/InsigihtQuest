from __future__ import annotations

from datetime import date, datetime, timezone
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from insightquest import exports


def _sheets() -> dict[str, pd.DataFrame]:
    return {
        "Sales": pd.DataFrame({"name": ["=cmd", "Zoé"], "qty": [1, 2]}),
        "Other": pd.DataFrame({"x": [1, 2, 3]}),
    }


def test_excel_without_metadata_has_no_filters_sheet():
    workbook = load_workbook(BytesIO(exports.build_excel_workbook(_sheets())))
    assert workbook.sheetnames == ["Sales", "Other"]
    assert workbook["Sales"]["A2"].value == "=cmd"
    assert workbook["Sales"]["A2"].data_type == "s"


def test_excel_filters_sheet_round_trip():
    payload = exports.build_excel_workbook(
        _sheets(),
        page="Inventory risk",
        filters={"Territory": ["QC", "ON"], "Max rows": 5000},
        date_range=(date(2026, 1, 1), date(2026, 3, 31)),
        generated_at=datetime(2026, 3, 4, 5, 6, 7, tzinfo=timezone.utc),
        proxy_label="PurchaseDate is a warranty-start proxy, not a confirmed shipping date",
        language="en",
    )
    workbook = load_workbook(BytesIO(payload))
    assert workbook.sheetnames == ["Sales", "Other", "Filters"]
    rows = {row[0]: row[1] for row in workbook["Filters"].iter_rows(min_row=2, values_only=True)}
    assert rows["Page"] == "Inventory risk"
    assert rows["Territory"] == "QC, ON"
    assert rows["Max rows"] == "5000"
    assert rows["Date range"] == "2026-01-01 – 2026-03-31"
    assert rows["Generated at (UTC)"] == "2026-03-04 05:06:07"
    assert rows["Row count"] == "5"
    assert "warranty-start proxy" in rows["Date note"]


def test_filters_sheet_localised_and_name_collision():
    sheets = {"Filtres": pd.DataFrame({"a": [1]})}
    workbook = load_workbook(BytesIO(exports.build_excel_workbook(sheets, page="p", language="fr_CA")))
    assert workbook.sheetnames == ["Filtres (2)", "Filtres"]
    assert workbook["Filtres"]["A1"].value == "Champ"


def test_filename_convention():
    stamp = datetime(2026, 3, 4, tzinfo=timezone.utc)
    assert exports.build_export_filename("Inventory Risk", today=stamp) == "insightquest_inventory_risk_20260304.xlsx"
    assert (
        exports.build_export_filename("Écarts / Territoire", "csv", table="Top 10", today=stamp)
        == "insightquest_ecarts_territoire_top_10_20260304.csv"
    )


def test_csv_round_trip_and_formula_protection():
    frame = pd.DataFrame({"name": ["=cmd", "Zoé"], "qty": [1.5, float("inf")]})
    payload = exports.build_csv(frame)
    assert payload.startswith(b"\xef\xbb\xbf")
    parsed = pd.read_csv(BytesIO(payload), encoding="utf-8-sig")
    assert parsed["name"].tolist() == ["'=cmd", "Zoé"]
    assert parsed["qty"].iloc[0] == 1.5
    assert pd.isna(parsed["qty"].iloc[1])
